from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DOCX = ROOT / "public" / "doc" / "芝麻日历_软件著作权源代码文档_前端前30页.docx"
SOFTWARE_LABEL = "芝麻日历 [V1.0.0]"
SOURCE_RANGES = [
    ("pages/calendar/index.vue", 1, 500),
    ("pages/tasks/index.vue", 1, 500),
    ("pages/rewards/index.vue", 1, 500),
]


def add_field(paragraph, field_name):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_name
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def set_run_font(run, name="Courier New", size=Pt(9)):
    run.font.name = name
    run.font.size = size
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_paragraph_spacing(paragraph, line_pt=14.0):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(line_pt)
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), str(int(line_pt * 20)))
    spacing.set(qn("w:lineRule"), "exact")


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(0.8)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.4)

    header = section.header
    header.is_linked_to_previous = False
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_para.text = ""
    header_run = header_para.add_run(SOFTWARE_LABEL)
    set_run_font(header_run, "SimSun", Pt(9))

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.text = ""
    add_field(footer_para, "PAGE")
    for run in footer_para.runs:
        set_run_font(run, "Courier New", Pt(9))


def build_document():
    source_lines = []
    for file_name, start, end in SOURCE_RANGES:
        path = ROOT / file_name
        lines = path.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
        effective_lines = [line for line in lines if line.strip()]
        source_lines.extend(effective_lines[start - 1 : end])

    if len(source_lines) != 1500:
        raise ValueError(f"Expected 1500 source lines, got {len(source_lines)}")

    raw_pages = [source_lines[index : index + 50] for index in range(0, 1500, 50)]

    doc = Document()
    configure_section(doc.sections[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Courier New"
    normal.font.size = Pt(9)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")

    for index, page in enumerate(raw_pages):
        if index:
            doc.add_section(WD_SECTION_START.NEW_PAGE)
            configure_section(doc.sections[-1])

        if len(page) != 50:
            raise ValueError(f"Page {index + 1} expected 50 source lines, got {len(page)}")

        for line in page:
            paragraph = doc.add_paragraph()
            set_paragraph_spacing(paragraph, 12.8)
            run = paragraph.add_run(line)
            set_run_font(run, "Courier New", Pt(9))

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_document()
